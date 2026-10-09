import re, sys, unicodedata
from docx import Document
from bs4 import BeautifulSoup
def norm(s):
    s = unicodedata.normalize("NFKC", s)
    for a,b in [("\u2019","'"),("\u2018","'"),("\u201c",'"'),("\u201d",'"'),("\u2013","-"),("\u2014","-"),("\u00a0"," "),("\u2022"," "),("\u2192"," "),("\u2190"," ")]:
        s = s.replace(a,b)
    s = re.sub(r"[^\w%$.,;:'\"()/&+-]+"," ",s)
    return re.sub(r"\s+"," ",s).strip().lower()
soup = BeautifulSoup(open(sys.argv[2]).read(),"html.parser")
for t in soup.find_all(["b","em","strong","i","span"]): t.unwrap()
soup.smooth()
H = norm(soup.get_text(" "))
d = Document(sys.argv[1])
def ptext(p): return "".join(r.text for r in p.runs if not r.font.superscript)
segs=[]
for p in d.paragraphs: segs.append(ptext(p))
for t in d.tables:
    for r in t.rows:
        seen=set()
        for c in r.cells:
            if id(c._tc) in seen: continue
            seen.add(id(c._tc)); segs += [ptext(p) for p in c.paragraphs]
LABELS=("what each party books today:","the question it raises:","where the framework stops:","practical perspective:")
total=0; miss=[]
for s in segs:
    for part in re.split(r"(?<=[.?!])\s+|\n", s):
        n = norm(part)
        if len(n) < 22: continue
        total += 1
        if n in H: continue
        m = n
        for L in LABELS:
            if m.startswith(L): m = m[len(L):].strip()
        if m in H: continue
        if "  " in part:
            a,b = part.split("  ",1)
            if norm(a) in H and norm(b) in H: continue
        miss.append(part.strip())
print("sentences checked:", total, "| not found:", len(miss))
for m in miss: print(" -", m[:170])
